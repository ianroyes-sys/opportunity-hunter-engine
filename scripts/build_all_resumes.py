#!/usr/bin/env python3
"""Build every resume variant from one structured content file.

Content lives in profile/resume.json (copy profile/resume.example.json to start).
Each variant picks which entries and bullets to show, in what order, so a tech,
finance or operations resume all come from the same facts.

Outputs:
  resumes/upload/<variant>/<File_Stem>.pdf   <- what the apply lane uploads
  resumes/generated/resume_<variant>.{md,pdf}
  resumes/<File_Stem>_<variant>.{pdf,txt,docx}  for variants listed in "docxVariants"
  profile/resume.txt                            plain text of "defaultVariant"

Every PDF must stay exactly one page; the build asserts it.
Usage: python3 scripts/build_all_resumes.py [--content profile/resume.json]
"""

import argparse
import json
from pathlib import Path
import re
import shutil
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]

# Populated by load(); module-level so the renderers below stay simple.
NAME = CONTACT_1 = CONTACT_2 = EDU_NOTE = LEADERSHIP = FILE_STEM = ""
LINKS, SKILLS, VARIANTS = {}, {}, {}
EDU_HEAD = ("", "", "", "")
CONFIG = {}


def load(path):
    global NAME, CONTACT_1, CONTACT_2, EDU_NOTE, LEADERSHIP, FILE_STEM, LINKS, SKILLS, VARIANTS, EDU_HEAD, CONFIG
    c = json.loads(Path(path).read_text())
    CONFIG = c
    NAME = c["name"]
    CONTACT_1, CONTACT_2 = c["contact"]
    LINKS = c.get("links", {})
    FILE_STEM = c.get("fileStem") or re.sub(r"[^A-Za-z0-9]+", "_", NAME).strip("_") + "_Resume"
    edu = c["education"]
    EDU_HEAD, EDU_NOTE, LEADERSHIP = tuple(edu["head"]), edu.get("note", ""), edu.get("extra", "")
    SKILLS = {k: [tuple(x) for x in v] for k, v in c["skills"].items()}
    entries = c["entries"]

    def entry(ref):
        eid, keys = ref
        e = entries[eid]
        return {"head": tuple(e["head"]), "bullets": [e["bullets"][k] for k in keys]}

    VARIANTS = {}
    for key, v in c["variants"].items():
        sections = []
        for title, content in v["sections"]:
            if title in ("Skills", "Education"):
                sections.append((title, content))
            else:
                sections.append((title, [entry(r) for r in content]))
        VARIANTS[key] = {"headline": v["headline"], "summary": v["summary"], "sections": sections}


# ---------------------------------------------------------------- PDF
INK, ACCENT, MUTED = HexColor("#111827"), HexColor("#1f3a5f"), HexColor("#4b5563")


def styles(scale):
    f = lambda x: x * scale
    return {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=f(21), leading=f(24), textColor=INK, alignment=1),
        "headline": ParagraphStyle("hl", fontName="Helvetica-Bold", fontSize=f(10.2), leading=f(13), textColor=ACCENT, alignment=1, spaceBefore=1),
        "contact": ParagraphStyle("c", fontName="Helvetica", fontSize=f(8.6), leading=f(11.2), textColor=MUTED, alignment=1),
        "summary": ParagraphStyle("s", fontName="Helvetica", fontSize=f(9.4), leading=f(12.4), textColor=INK, spaceBefore=f(5)),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=f(10), leading=f(12), textColor=ACCENT, spaceBefore=f(7), spaceAfter=1),
        "org": ParagraphStyle("o", fontName="Helvetica-Bold", fontSize=f(9.6), leading=f(12), textColor=INK),
        "date": ParagraphStyle("d", fontName="Helvetica", fontSize=f(9.2), leading=f(12), textColor=INK, alignment=2),
        "role": ParagraphStyle("r", fontName="Helvetica-Oblique", fontSize=f(9.2), leading=f(11.6), textColor=INK),
        "loc": ParagraphStyle("l", fontName="Helvetica-Oblique", fontSize=f(9), leading=f(11.6), textColor=MUTED, alignment=2),
        "li": ParagraphStyle("li", fontName="Helvetica", fontSize=f(9.2), leading=f(11.9), leftIndent=11, firstLineIndent=0,
                             bulletIndent=2, textColor=INK, spaceAfter=f(1)),
        "skill": ParagraphStyle("sk", fontName="Helvetica", fontSize=f(9.2), leading=f(12), textColor=INK, spaceAfter=f(1)),
    }


def linkify(text):
    labels = sorted(LINKS, key=len, reverse=True)
    pattern = re.compile("|".join(re.escape(label) for label in labels))
    return pattern.sub(lambda m: f'<link href="{LINKS[m.group(0)]}" color="#1f3a5f">{m.group(0)}</link>', escape(text))


def two_col(left, right, ls, rs, width):
    t = Table([[Paragraph(left, ls), Paragraph(right, rs)]], colWidths=[width * 0.74, width * 0.26])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    return t


def story_for(v, s, width):
    story = [Paragraph(NAME, s["name"]), Paragraph(v["headline"], s["headline"]),
             Paragraph(escape(CONTACT_1), s["contact"]), Paragraph(linkify(CONTACT_2), s["contact"]),
             Paragraph(escape(v["summary"]), s["summary"])]

    def h2(title):
        story.append(Paragraph(title.upper(), s["h2"]))
        story.append(HRFlowable(width="100%", thickness=0.6, color=ACCENT, spaceBefore=0, spaceAfter=3))

    def add_entry(e):
        org, date, role, loc = e["head"]
        story.append(Spacer(1, 3))
        story.append(two_col(linkify(org), escape(date), s["org"], s["date"], width))
        if role or loc:
            story.append(two_col(escape(role), escape(loc), s["role"], s["loc"], width))
        for b in e["bullets"]:
            story.append(Paragraph(escape(b), s["li"], bulletText="•"))

    for title, content in v["sections"]:
        h2(title)
        if title == "Skills":
            for label, items in SKILLS[content]:
                story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(items)}", s["skill"]))
        elif title == "Education":
            org, date, role, loc = EDU_HEAD
            story.append(two_col(org, date, s["org"], s["date"], width))
            story.append(two_col(escape(role), loc, s["role"], s["loc"], width))
            story.append(Paragraph(f"<b>{escape(EDU_NOTE)}</b>", s["skill"]))
            story.append(Paragraph(escape(LEADERSHIP), s["skill"]))
        else:
            for e in content:
                add_entry(e)
    return story


def build_pdf(path, variant_key):
    v = VARIANTS[variant_key]
    margin = 0.5 * 72
    width = 612 - 2 * margin - 12  # SimpleDocTemplate frames pad 6pt per side
    # Fill the page: grow type until the next step would spill onto page 2.
    best = None
    for scale in [1.12, 1.09, 1.06, 1.03, 1.0, 0.97, 0.94, 0.91, 0.88]:
        doc = SimpleDocTemplate(str(path), pagesize=(612, 792), leftMargin=margin, rightMargin=margin,
                                topMargin=0.42 * 72, bottomMargin=0.42 * 72,
                                title=f"{NAME} — Resume", author=NAME, subject=v["headline"])
        doc.build(story_for(v, styles(scale), width))
        if len(PdfReader(str(path)).pages) == 1:
            best = scale
            break
    assert best, f"{path.name}: cannot fit on one page"
    return best


# ---------------------------------------------------------------- text / markdown / docx
def plain_text(variant_key):
    v = VARIANTS[variant_key]
    lines = [NAME.upper(), v["headline"].replace("  ·  ", " | "), CONTACT_1.replace("  ·  ", " | "),
             CONTACT_2.replace("  ·  ", " | "), "", v["summary"], ""]
    for title, content in v["sections"]:
        lines.append(title.upper())
        if title == "Skills":
            lines += [f"{label}: {items}" for label, items in SKILLS[content]]
        elif title == "Education":
            org, date, role, loc = EDU_HEAD
            lines += [f"{org} — {role} | {loc} | {date}", EDU_NOTE, LEADERSHIP]
        else:
            for e in content:
                org, date, role, loc = e["head"]
                lines.append(" | ".join(x for x in [f"{org} — {role}", loc, date] if x))
                lines += [f"- {b}" for b in e["bullets"]]
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def markdown(variant_key):
    v = VARIANTS[variant_key]
    lines = [f"# {NAME}", f"**{v['headline']}**", "", CONTACT_1, CONTACT_2, "", v["summary"], ""]
    for title, content in v["sections"]:
        lines.append(f"## {title}")
        if title == "Skills":
            lines += [f"- **{label}:** {items}" for label, items in SKILLS[content]]
        elif title == "Education":
            org, date, role, loc = EDU_HEAD
            lines += [f"### {org} — {role} | {loc} | {date}", EDU_NOTE, "", LEADERSHIP]
        else:
            for e in content:
                org, date, role, loc = e["head"]
                lines.append("### " + " | ".join(x for x in [f"{org} — {role}", loc, date] if x))
                lines += [f"- {b}" for b in e["bullets"]]
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def build_docx(path, variant_key):
    v = VARIANTS[variant_key]
    doc = Document()
    sec = doc.sections[0]
    for side in ("top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.45))
    sec.left_margin = sec.right_margin = Inches(0.55)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(9.5)
    normal.paragraph_format.space_after = Pt(1)
    tab_pos = Inches(8.5 - 1.1)

    def para(text="", bold=False, italic=False, size=None, center=False, color=None, right=None):
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if right is not None:
            p.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
            rr = p.add_run("\t" + right)
            rr.bold, rr.italic = bold, italic
        return p

    para(NAME, bold=True, size=18, center=True)
    para(v["headline"], bold=True, size=10, center=True, color="1F3A5F")
    para(CONTACT_1, size=8.5, center=True)
    para(CONTACT_2, size=8.5, center=True)
    para(v["summary"])
    for title, content in v["sections"]:
        h = para(title.upper(), bold=True, size=10, color="1F3A5F")
        h.paragraph_format.space_before = Pt(6)
        if title == "Skills":
            for label, items in SKILLS[content]:
                p = doc.add_paragraph()
                p.add_run(label + ": ").bold = True
                p.add_run(items)
        elif title == "Education":
            org, date, role, loc = EDU_HEAD
            para(org, bold=True, right=date)
            para(role, italic=True, right=loc)
            para(EDU_NOTE, bold=True)
            para(LEADERSHIP)
        else:
            for e in content:
                org, date, role, loc = e["head"]
                para(org, bold=True, right=date)
                para(role, italic=True, right=loc)
                for b in e["bullets"]:
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(b)
    doc.save(path)


# ---------------------------------------------------------------- main
def copy(src, dst):
    dst = Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)


def main():
    ap = argparse.ArgumentParser(description="Build one-page resume variants from structured content.")
    default = ROOT / "profile" / "resume.json"
    if not default.exists():
        default = ROOT / "profile" / "resume.example.json"
    ap.add_argument("--content", type=Path, default=default)
    load(ap.parse_args().content)

    gen = ROOT / "resumes" / "generated"
    gen.mkdir(parents=True, exist_ok=True)
    report = []
    for key in VARIANTS:
        pdf = gen / f"resume_{key}.pdf"
        scale = build_pdf(pdf, key)
        (gen / f"resume_{key}.md").write_text(markdown(key))
        copy(pdf, ROOT / "resumes" / "upload" / key / f"{FILE_STEM}.pdf")
        report.append(f"{key}: 1 page (type scale {scale})")

    for key in CONFIG.get("docxVariants", []):
        stem = ROOT / "resumes" / f"{FILE_STEM}_{key}"
        copy(gen / f"resume_{key}.pdf", stem.with_suffix(".pdf"))
        stem.with_suffix(".txt").write_text(plain_text(key))
        build_docx(stem.with_suffix(".docx"), key)

    default_variant = CONFIG.get("defaultVariant")
    if default_variant:
        (ROOT / "profile" / "resume.txt").write_text(plain_text(default_variant))

    print("\n".join(report))


if __name__ == "__main__":
    main()
